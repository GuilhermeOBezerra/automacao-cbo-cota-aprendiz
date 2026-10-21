import os
import time
import glob
import shutil

import openpyxl
import tkinter as tk
from tkinter import filedialog, messagebox

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (TimeoutException, NoSuchElementException,
                                        StaleElementReferenceException, WebDriverException)

URL_BUSCA = "https://www.mtecbo.gov.br/cbosite/pages/pesquisas/BuscaPorCodigo.jsf"
TEMPO_MAX_CAPTCHA = 600      # segundos (10 min) aguardando a resolução manual
USAR_PROXY_MANUAL = False    # True apenas em redes que exigem login de proxy no navegador

LOC_CAMPO_CBO = (By.CSS_SELECTOR,
                 "input.cbo2002Format, input[name^='formBuscaPorCodigo:'][type='text']")

# ================================================================ CAPTCHA
SELETORES_CAPTCHA = [
    "iframe[src*='recaptcha']",
    "iframe[src*='hcaptcha']",
    "iframe[src*='turnstile']",
    "iframe[title*='captcha' i]",
    "div.g-recaptcha", "div.h-captcha", "div.cf-turnstile",
    "img[src*='captcha' i]", "img[id*='captcha' i]", "img[alt*='captcha' i]",
    "input[id*='captcha' i]", "input[name*='captcha' i]",
    "[id*='captcha' i]", "[class*='captcha' i]",
]
TEXTOS_CAPTCHA = ["não sou um robô", "nao sou um robo", "i'm not a robot",
                  "verify you are human", "just a moment", "digite os caracteres"]


def _visivel(el):
    try:
        if not el.is_displayed():
            return False
        tam = el.size
        if tam.get("width", 0) < 5 or tam.get("height", 0) < 5:
            return False
        src = (el.get_attribute("src") or "").lower()
        if "size=invisible" in src:  # selo do reCAPTCHA invisível: ignora
            return False
        return True
    except (StaleElementReferenceException, WebDriverException):
        return False


def captcha_presente(nav):
    """Retorna True se houver algum captcha visível na página."""
    try:
        for sel in SELETORES_CAPTCHA:
            if any(_visivel(e) for e in nav.find_elements(By.CSS_SELECTOR, sel)):
                return True
        corpo = nav.find_element(By.TAG_NAME, "body").text.lower()
        titulo = (nav.title or "").lower()
        return any(t in corpo or t in titulo for t in TEXTOS_CAPTCHA)
    except WebDriverException:
        return False


def token_captcha_resolvido(nav):
    """True quando o token do reCAPTCHA/hCaptcha/Turnstile já foi gerado."""
    js = """
    const sel = ['textarea[name="g-recaptcha-response"]',
                 'textarea[name="h-captcha-response"]',
                 'input[name="cf-turnstile-response"]'];
    for (const s of sel) for (const e of document.querySelectorAll(s))
        if (e.value && e.value.length > 20) return true;
    return false;
    """
    try:
        return bool(nav.execute_script(js))
    except WebDriverException:
        return False


def alertar_usuario():
    print("\a")
    try:
        import winsound
        winsound.MessageBeep()
    except Exception:
        pass


def aguardar_captcha(nav, condicao_sucesso=None, intervalo=1.0):
    """
    Se houver captcha, espera até que o token seja gerado, o captcha desapareça
    ou condicao_sucesso(nav) seja verdadeira. Retorna True se havia captcha.
    """
    if not captcha_presente(nav):
        return False

    print(">>> CAPTCHA detectado! Resolva no navegador "
          "(se for captcha de imagem, digite o texto e clique em Consultar).")
    alertar_usuario()
    limite = time.time() + TEMPO_MAX_CAPTCHA
    while time.time() < limite:
        try:
            if condicao_sucesso and condicao_sucesso(nav):
                break
        except WebDriverException:
            pass
        if token_captcha_resolvido(nav) or not captcha_presente(nav):
            break
        time.sleep(intervalo)
    else:
        raise TimeoutException("Captcha não resolvido no tempo limite.")

    print(">>> Captcha resolvido. Continuando...")
    time.sleep(1)
    return True


def esperar_elemento(nav, locator, timeout=10):
    """Espera um elemento existir, tratando captchas que surjam durante a espera."""
    limite = time.time() + timeout
    while True:
        try:
            return WebDriverWait(nav, 1).until(EC.presence_of_element_located(locator))
        except TimeoutException:
            if aguardar_captcha(nav, condicao_sucesso=lambda d: d.find_elements(*locator)):
                limite = time.time() + timeout
                continue
            if time.time() >= limite:
                raise


def clicar(nav, elemento):
    """Clique normal; se algo estiver cobrindo o elemento, clica via JavaScript."""
    try:
        nav.execute_script("arguments[0].scrollIntoView({block:'center'});", elemento)
        elemento.click()
    except WebDriverException:
        nav.execute_script("arguments[0].click();", elemento)


# ================================================================ CAMPO DA CBO
def localizar_campo_cbo(nav, timeout=20):
    """Procura o campo da CBO na página principal e dentro de iframes; trata captcha."""
    limite = time.time() + timeout
    while time.time() < limite:
        nav.switch_to.default_content()
        campos = [c for c in nav.find_elements(*LOC_CAMPO_CBO) if c.is_displayed()]
        if campos:
            return campos[0]

        for frame in nav.find_elements(By.TAG_NAME, "iframe"):
            try:
                nav.switch_to.frame(frame)
                campos = [c for c in nav.find_elements(*LOC_CAMPO_CBO) if c.is_displayed()]
                if campos:
                    print("Campo da CBO encontrado dentro de um iframe.")
                    return campos[0]
            except WebDriverException:
                pass
            nav.switch_to.default_content()

        nav.switch_to.default_content()
        if aguardar_captcha(nav, condicao_sucesso=lambda d: d.find_elements(*LOC_CAMPO_CBO)):
            limite = time.time() + timeout
            continue
        time.sleep(0.5)

    nav.switch_to.default_content()
    nav.save_screenshot("erro_campo_cbo.png")
    print(f"URL atual: {nav.current_url}")
    print(f"Título: {nav.title}")
    print(f"Iframes na página: {len(nav.find_elements(By.TAG_NAME, 'iframe'))}")
    raise TimeoutException("Campo da CBO não encontrado (veja erro_campo_cbo.png).")


def preencher_campo_cbo(nav, campo, codigo_cbo):
    """Digita o código respeitando a máscara 9999-99; usa JavaScript como alternativa."""
    cbo_formatado = f"{codigo_cbo[:4]}-{codigo_cbo[4:]}"
    nav.execute_script("arguments[0].scrollIntoView({block:'center'});", campo)
    try:
        campo.click()
    except WebDriverException:
        nav.execute_script("arguments[0].focus();", campo)
    campo.clear()
    for ch in codigo_cbo:
        campo.send_keys(ch)
        time.sleep(0.05)

    valor = (campo.get_attribute("value") or "").replace("-", "")
    if valor != codigo_cbo:
        print(f"Aviso: máscara recusou digitação (valor='{valor}'). Preenchendo via JS.")
        nav.execute_script("""
            const el = arguments[0];
            el.value = arguments[1];
            el.dispatchEvent(new Event('input',  {bubbles:true}));
            el.dispatchEvent(new Event('change', {bubbles:true}));
            el.dispatchEvent(new Event('blur',   {bubbles:true}));
        """, campo, cbo_formatado)


# ================================================================ UTILIDADES
def obter_caminho_excel():
    root = tk.Tk()
    root.withdraw()
    caminho = filedialog.askopenfilename(
        title="Selecione o arquivo Excel",
        filetypes=[("Arquivos Excel", "*.xlsx")]
    )
    if caminho:
        print(f"Arquivo selecionado: {os.path.basename(caminho)}")
    return caminho


def gerar_backup(caminho_excel, max_backups=5):
    pasta_backup = os.path.join(os.path.dirname(caminho_excel), "backup")
    os.makedirs(pasta_backup, exist_ok=True)

    backups = sorted(glob.glob(os.path.join(pasta_backup, "Backup CBOs*.xlsx")),
                     key=os.path.getmtime)
    numeros = []
    for f in backups:
        try:
            numeros.append(int(os.path.basename(f).replace("Backup CBOs", "").replace(".xlsx", "")))
        except ValueError:
            pass

    if len(backups) >= max_backups:
        os.remove(backups[0])
        print(f"Backup antigo removido: '{os.path.basename(backups[0])}'")

    proximo = (max(numeros) if numeros else 0) + 1
    nome_backup = os.path.join(pasta_backup, f"Backup CBOs{proximo}.xlsx")
    shutil.copy2(caminho_excel, nome_backup)
    print(f"Backup criado: '{os.path.basename(nome_backup)}'")


def normalizar_cbo(valor):
    """Converte 253215, 253215.0 ou '2532-15' em '253215'."""
    if isinstance(valor, (int, float)):
        return str(int(valor)).zfill(6)
    return "".join(ch for ch in str(valor) if ch.isdigit()).zfill(6)


# ================================================================ PESQUISA
def extrair_info_cbo(codigo_cbo, nav):
    cbo_formatado = f"{codigo_cbo[:4]}-{codigo_cbo[4:]}"
    loc_link = (By.LINK_TEXT, cbo_formatado)
    loc_botao = (By.ID, "formBuscaPorCodigo:btConsultarCodigo")
    print(f"Buscando CBO: {cbo_formatado}")

    # 1) Campo de busca
    campo = localizar_campo_cbo(nav)
    preencher_campo_cbo(nav, campo, codigo_cbo)
    time.sleep(0.5)

    # 2) Captcha antes de consultar
    aguardar_captcha(nav, condicao_sucesso=lambda d: d.find_elements(*loc_link))

    # 3) Consultar (se o usuário ainda não consultou ao resolver o captcha)
    if not nav.find_elements(*loc_link):
        clicar(nav, esperar_elemento(nav, loc_botao))
        time.sleep(1)

    # 4) Resultado
    clicar(nav, esperar_elemento(nav, loc_link, timeout=15))
    time.sleep(1)
    nav.switch_to.default_content()

    # 5) Menu do relatório completo
    menu = esperar_elemento(nav, (By.XPATH, '//*[@id="menu_completo_dataprev"]/li[9]/a'), timeout=15)
    nav.execute_script("arguments[0].click();", menu)
    time.sleep(1)
    aguardar_captcha(nav)

    familia = codigo_cbo[:4]
    try:
        td_codigo = esperar_elemento(nav, (By.XPATH, f"//td[normalize-space(text())='{familia}']"))
        titulo = td_codigo.find_element(By.XPATH, "./following-sibling::td[1]").text
    except (NoSuchElementException, TimeoutException):
        print("Aviso: título não achado pelo código. Usando XPath alternativo.")
        titulo = esperar_elemento(nav, (By.XPATH, '//*[@id="formSite022"]/table[1]/tbody/tr/td[2]')).text

    # A página já carregou (o título foi encontrado): basta conferir o texto uma vez
    trecho = "artigo 429 da Consolidação das Leis do Trabalho - CLT"
    corpo = nav.find_element(By.TAG_NAME, "body").text
    tem_clt = "Sim" if trecho in corpo else "Não"

    # 6) Volta para a busca
    nav.get(URL_BUSCA)
    try:
        localizar_campo_cbo(nav)
    except TimeoutException:
        print("Erro: página de pesquisa não recarregou.")

    return titulo, tem_clt


# ================================================================ MAIN
def main():
    start = time.time()
    nav = None
    sucesso = False

    try:
        caminho_excel = obter_caminho_excel()
        if not caminho_excel:
            messagebox.showinfo("Aviso", "Nenhum arquivo selecionado.")
            return

        gerar_backup(caminho_excel)

        # Lê e escreve na MESMA aba com openpyxl
        wb = openpyxl.load_workbook(caminho_excel)
        sheet = wb.active
        colunas = {str(c.value).strip().upper(): i
                   for i, c in enumerate(sheet[1], 1) if c.value}
        if not all(c in colunas for c in ["CBO", "TÍTULO", "COTA"]):
            messagebox.showerror("Erro de Coluna", "Faltam as colunas CBO, TÍTULO ou COTA.")
            return
        col_cbo, col_titulo, col_cota = colunas["CBO"], colunas["TÍTULO"], colunas["COTA"]

        # Retoma da primeira linha com COTA vazia
        inicio = next((r for r in range(2, sheet.max_row + 1)
                       if sheet.cell(r, col_cota).value in (None, "")), None)
        if inicio is None:
            messagebox.showinfo("Concluído", "A coluna COTA já está toda preenchida.")
            return
        print(f"Iniciando na linha {inicio}.")

        nav = webdriver.Chrome()
        if USAR_PROXY_MANUAL:
            nav.get("https://www.google.com")
            input("Faça o login do proxy no navegador e pressione Enter para continuar...")

        nav.get(URL_BUSCA)
        try:
            WebDriverWait(nav, 7).until(EC.presence_of_element_located((By.ID, "logo")))
        except TimeoutException:
            # O Chrome pode exibir "Sua conexão não é particular" (certificado do site gov.br)
            try:
                WebDriverWait(nav, 3).until(EC.element_to_be_clickable((By.ID, "details-button"))).click()
                WebDriverWait(nav, 3).until(EC.element_to_be_clickable((By.ID, "proceed-link"))).click()
                print("Aviso de certificado do Chrome ignorado para o site do MTE.")
            except (TimeoutException, NoSuchElementException):
                print("Não foi possível passar pelo aviso de certificado.")
        aguardar_captcha(nav)

        cache = {}
        for linha in range(inicio, sheet.max_row + 1):
            cbo = sheet.cell(linha, col_cbo).value
            if cbo in (None, ""):
                print("CBO vazia. Fim da lista.")
                break

            codigo = normalizar_cbo(cbo)
            familia = codigo[:4]
            if familia in cache:
                titulo, tem_clt = cache[familia]
                print(f"CBO {codigo}: família {familia} em cache.")
            else:
                try:
                    titulo, tem_clt = extrair_info_cbo(codigo, nav)
                    cache[familia] = (titulo, tem_clt)
                except Exception as e:
                    print(f"Erro na CBO {codigo}: {e}")
                    messagebox.showerror("Erro", "Erro encontrado. Progresso salvo; o script será encerrado.")
                    return

            sheet.cell(linha, col_titulo).value = titulo
            sheet.cell(linha, col_cota).value = tem_clt
            try:
                wb.save(caminho_excel)
            except PermissionError:
                messagebox.showerror("Arquivo aberto",
                                     "Feche a planilha no Excel e rode o script novamente.")
                return
            print(f"Linha {linha}: CBO {codigo} -> TÍTULO: '{titulo}', COTA: '{tem_clt}' (salvo)")

        sucesso = True

    finally:
        if nav:
            nav.quit()
        dur = time.time() - start
        tempo = f"{int(dur // 60)} min {int(dur % 60)} s" if dur >= 60 else f"{dur:.2f} s"
        if sucesso:
            messagebox.showinfo("Concluído", f"Script concluído com sucesso em {tempo}.")
        else:
            print(f"Script encerrado após {tempo}.")
        input("Pressione Enter para sair...")


if __name__ == "__main__":
    main()
